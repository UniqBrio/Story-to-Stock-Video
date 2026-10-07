#!/usr/bin/env python3
"""
story_reader.py  (SVOS v2 render layer)
───────────────
Reads story_plan.xlsx and writes project.json. Every other stage reads
project.json — Excel is only touched here (and by the fetcher's log writer).

What's new in v2
  • Columns are located by HEADER NAME (row 3), so optional columns can be
    added in any order: Treatment · Text Style · Text Anim · Text Color ·
    Card BG · Logo Bug · Font Size. Old sheets still work unchanged.
  • Shot IDs are zero-padded ("1" → "001").
  • Validation: trim window vs duration, text window inside the shot, pre-cut
    safety, enum values, missing files, Reels length, restraint budget (README §2).
  • API keys can come from env vars (PEXELS_API_KEY, PIXABAY_API_KEY,
    UNSPLASH_API_KEY) so they never need to live in the sheet.

Usage:
    python story_reader.py story_plan.xlsx
    python story_reader.py story_plan.xlsx --out custom_project.json
    python story_reader.py story_plan.xlsx --strict      # warnings become errors
    python story_reader.py story_plan.xlsx --fresh       # forget previously fetched assets
"""

import sys
import json
import argparse
from pathlib import Path

from content_safety import subject_needs_illustration, check_text, PASS as SAFE_PASS, FAIL as SAFE_FAIL
from svos_common import (BRAND, normalise_shot_id, shot_kind, shot_treatment, words,
                         env_key, to_float, to_int, to_bool, hex_clean, compute_timeline,
                         timeline_total, REAL_TRANSITIONS, best_logo)

try:
    import openpyxl
except ImportError:
    sys.exit("❌  Run: pip install openpyxl")


# ── Header aliases → field names (lower-cased, punctuation-insensitive) ───────
HEADER_ALIASES = {
    "shot id":            "shot_id",
    "scene description":  "scene_desc",
    "duration (s)":       "duration",
    "duration":           "duration",
    "search keywords":    "keywords",
    "keywords":           "keywords",
    "asset priority":     "asset_priority",
    "asset type":         "asset_type",
    "source":             "source",
    "local file path":    "local_file",
    "local file":         "local_file",
    "trim in (s)":        "trim_in",
    "trim out (s)":       "trim_out",
    "color grade":        "color_grade",
    "colour grade":       "color_grade",
    "ken burns":          "ken_burns",
    "text overlay":       "text_overlay",
    "text start (s)":     "text_start",
    "text end (s)":       "text_end",
    "text position":      "text_position",
    "transition out":     "transition_out",
    "trans dur (s)":      "trans_dur",
    "vo file path":       "vo_file",
    "status":             "status",
    # v2 optional columns
    "treatment":          "treatment",
    "text style":         "text_style",
    "text anim":          "text_anim",
    "text color":         "text_color",
    "text colour":        "text_color",
    "card bg":            "card_bg",
    "logo bug":           "logo_bug",
    "font size":          "text_size",
    "vo line":            "vo_line",
    "visual style":       "visual_style",
    "notes":              "notes",
}

# Positional fallback (legacy sheets without a readable header row)
LEGACY_COLS = ["shot_id", "scene_desc", "duration", "keywords", "asset_priority", "asset_type",
               "source", "local_file", "trim_in", "trim_out", "color_grade", "ken_burns",
               "text_overlay", "text_start", "text_end", "text_position", "transition_out",
               "trans_dur", "vo_file", "status"]

ALLOWED = {
    "asset_priority": {"video_first", "image_first", "video_only", "image_only"},
    "color_grade":    {"chaos", "pivot", "cta", "none"},
    "ken_burns":      {"zoom_in", "zoom_out", "pan_left", "pan_right", "none"},
    "text_position":  {"center", "lower_left", "bottom_center", "top_center", "upper_third", "lower_third"},
    "transition_out": {"cut", "dissolve", "fade_black", "fade_white", "none"},
    "treatment":      {"T1", "T2", "T3", "T4", "T5", ""},
    "text_style":     {"auto", "caption", "keyword", "card", "cta", ""},
    "text_anim":      {"rise", "fade", "pop", "none", ""},
    "logo_bug":       {"yes", "no", "auto", ""},
    "visual_style":   {"auto", "photo", "illustration", ""},
}

DEFAULT_SETTINGS = {
    # existing
    "project_name": "my_video", "output_filename": "final_video.mp4",
    "output_width": "1080", "output_height": "1920", "fps": "30",
    "music_volume": "0.8", "music_fade_in": "0.5", "music_fade_out": "2.0",
    "bgm_volume": "0.4", "vo_volume": "1.0",
    "logo_position": "bottom_right", "logo_scale": "0.15", "logo_opacity": "0.9",
    "crf": "18", "preset": "slow", "text_font": "Inter", "text_size_default": "72", "text_shadow": "true",
    # v2 design + audio settings
    "brand_color":        BRAND["orange"],
    "card_bg_default":    BRAND["offwhite"],
    "card_text_color":    BRAND["near_black"],
    "logo_card_bg":       BRAND["purple"],
    "logo_card_scale":    "0.42",
    "logo_mode":          "end_only",      # end_only | bug | both | none
    "fonts_folder":       "",
    "text_font_tamil":    "",
    "text_anim_default":  "rise",
    "text_box":           "auto",          # auto | always | never  (semi-transparent plate behind captions)
    "unified_grade":      "warm_soft",     # none | warm_soft | clean_neutral | cool_calm
    "film_grain":         "false",
    "platform":           "reels",
    "safe_zones":         "true",
    "target_lufs":        "-14",
    "true_peak":          "-1.5",
    "duck_db":            "8",
    "sting_path":         "",
    "sting_volume":       "0.9",
    "vo_start":           "0.0",
    "overlay_budget":     "5",
    "card_budget":        "3",
    "cover_time":         "",
    "max_duration":       "90",
    # stock fetch (asset_fetcher)
    "max_queries_per_shot":  "3",
    "fetch_budget_pexels":   "150",     # requests per run — Pexels allows 200/hour
    "fetch_budget_pixabay":  "300",
    "fetch_budget_unsplash": "40",      # Unsplash demo keys allow 50/hour
    "fetch_cache_hours":     "24",
    "max_download_mb":       "300",
    # content safety (content_safety.py)
    "vo_language":           "auto",    # auto | en | ta — language(s) Whisper listens for
    "captions_path":         "",        # SRT / VTT / TXT captions to validate
    "illustration_only_subjects": "",   # comma list; blank = built-in swimwear list
    "safety_models_dir":     "",
    "safety_frames_per_second": "1",
}

# Fields the fetcher writes per shot. They live only in project.json, so phase 1
# carries them forward instead of forgetting every fetched / approved asset.
CARRY_FIELDS = ("local_file", "asset_type", "source", "status", "asset_id", "asset_url", "page_url", "author",
                "license", "asset_score", "asset_resolution", "asset_duration", "query_used", "approved")


def _norm_header(s) -> str:
    return " ".join(str(s or "").strip().lower().replace("–", "-").split())


def _clean_path(v, default: str = "") -> str:
    s = str(v).strip() if v not in (None, "") else default
    if s.startswith("←"):                            # old bgm_prompt_generator placeholder, not a path
        return default
    if len(s) >= 2 and s[0] == s[-1] and s[0] in ('"', "'"):
        s = s[1:-1].strip()
    return s


def _str(v, default: str = "") -> str:
    return str(v).strip() if v not in (None, "") else default


# ── Sheet readers ─────────────────────────────────────────────────────────────
def read_settings(wb) -> dict:
    ws = wb["Project Settings"]
    settings = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row and row[0]:
            key = str(row[0]).strip()
            val = row[1] if len(row) > 1 else None
            settings[key] = str(val).strip() if val is not None else ""
    return settings


def locate_columns(ws) -> tuple[dict, int]:
    """Return ({field: col_index}, first_data_row). Scans the first 6 rows for a header."""
    for r in range(1, 7):
        found = {}
        for c in range(1, ws.max_column + 1):
            h = _norm_header(ws.cell(row=r, column=c).value)
            if h in HEADER_ALIASES:
                found[HEADER_ALIASES[h]] = c
        if "shot_id" in found and "scene_desc" in found:
            return found, r + 1
    return {f: i + 1 for i, f in enumerate(LEGACY_COLS)}, 4


def read_shots(wb, subjects: list | None = None) -> tuple[list, list, list]:
    ws = wb["Shot Plan"]
    cols, first_row = locate_columns(ws)
    shots, warnings, errors = [], [], []

    def val(row, field, default=""):
        c = cols.get(field)
        if not c:
            return default
        v = ws.cell(row=row, column=c).value
        return v if v is not None else default

    seen_ids = set()
    for r in range(first_row, ws.max_row + 1):
        raw_id = val(r, "shot_id")
        if raw_id in ("", None):
            continue
        sid = normalise_shot_id(raw_id)
        if sid in seen_ids:
            errors.append(f"Shot {sid} — duplicate Shot ID")
        seen_ids.add(sid)

        duration = to_float(val(r, "duration"), 0.0)
        shot = {
            "shot_id":        sid,
            "scene_desc":     _str(val(r, "scene_desc")),
            "duration":       duration if duration > 0 else 4.0,
            "keywords":       [k.strip() for k in _str(val(r, "keywords")).split(",") if k.strip()],
            "asset_priority": _str(val(r, "asset_priority"), "video_first"),
            "asset_type":     _str(val(r, "asset_type")).lower(),
            "source":         _str(val(r, "source")),
            "local_file":     _clean_path(val(r, "local_file")),
            "trim_in":        to_float(val(r, "trim_in"), 0.0),
            "trim_out":       to_float(val(r, "trim_out"), 0.0),
            "color_grade":    _str(val(r, "color_grade"), "none").lower(),
            "ken_burns":      _str(val(r, "ken_burns"), "none").lower(),
            "text_overlay":   _str(val(r, "text_overlay")),
            "text_start":     to_float(val(r, "text_start"), 0.0),
            "text_end":       to_float(val(r, "text_end"), 0.0),
            "text_position":  _str(val(r, "text_position"), "center").lower(),
            "transition_out": _str(val(r, "transition_out"), "cut").lower(),
            "trans_dur":      to_float(val(r, "trans_dur"), 0.0),
            "vo_file":        _clean_path(val(r, "vo_file")),
            "status":         _str(val(r, "status"), "pending").lower(),
            # v2
            "treatment":      _str(val(r, "treatment")).upper(),
            "text_style":     _str(val(r, "text_style"), "auto").lower(),
            "text_anim":      _str(val(r, "text_anim")).lower(),
            "text_color":     _str(val(r, "text_color")),
            "card_bg":        _str(val(r, "card_bg")),
            "logo_bug":       _str(val(r, "logo_bug"), "auto").lower(),
            "text_size":      to_int(val(r, "text_size"), 0),
            "vo_line":        _str(val(r, "vo_line")),
            "visual_style":   _str(val(r, "visual_style"), "auto").lower(),
            "notes":          _str(val(r, "notes")),
        }
        if duration <= 0:
            warnings.append(f"Shot {sid} — Duration missing/zero; defaulted to 4.0s")

        # Enum validation
        for field, allowed in ALLOWED.items():
            v = shot.get(field, "")
            if v and v not in allowed:
                warnings.append(f"Shot {sid} — '{v}' is not a valid {field}. Allowed: {sorted(a for a in allowed if a)}")

        # Treatment + kind
        kind = shot_kind(shot)
        shot["treatment"] = shot_treatment(shot)
        shot["kind"] = kind
        if kind in ("card", "logo_card"):
            shot["asset_type"] = kind
            shot["asset_priority"] = "image_only"
            shot["keywords"] = []          # never fetch stock for cards
        elif kind == "product":
            shot["asset_type"] = shot["asset_type"] or "video"
            shot["keywords"] = []          # product inserts come from the Demo Module Library, not stock

        # Swimwear-type subjects may only use cartoon / illustrated / animated visuals
        if kind == "stock":
            subj = subject_needs_illustration(
                [shot["scene_desc"], " ".join(shot["keywords"]), shot["text_overlay"], shot["vo_line"]], subjects)
            if subj or shot["visual_style"] == "illustration":
                shot["illustration_only"] = True
                shot["illustration_reason"] = (f"subject “{subj}” — cartoon/illustration only, no real people"
                                               if subj else "Visual Style = illustration")
                if subj and shot["visual_style"] == "photo":
                    warnings.append(f"Shot {sid} — Visual Style = photo ignored: the subject “{subj}” may only use "
                                    f"cartoon / illustrated / animated visuals")

        # Trim window must equal the planned duration — otherwise every later overlay drifts
        d = shot["duration"]
        if shot["trim_out"] <= 0.0:
            shot["trim_out"] = shot["trim_in"] + d
        elif abs((shot["trim_out"] - shot["trim_in"]) - d) > 0.05 and kind not in ("card", "logo_card"):
            warnings.append(f"Shot {sid} — trim window {shot['trim_in']}–{shot['trim_out']}s "
                            f"≠ duration {d}s; using trim_in + duration")
            shot["trim_out"] = shot["trim_in"] + d

        # Text window sanity (text-overlay-timing-director: clear 0.3s before a cut, 0.5s before a transition)
        if shot["text_overlay"]:
            pre = 0.5 if shot["transition_out"] in REAL_TRANSITIONS else 0.3
            if shot["text_end"] <= 0.0:
                shot["text_end"] = round(max(shot["text_start"] + 1.0, d - pre), 2)
                warnings.append(f"Shot {sid} — Text End blank; set to {shot['text_end']}s")
            if shot["text_end"] > d - pre + 0.001 and kind not in ("card", "logo_card"):
                new_end = round(max(shot["text_start"] + 0.8, d - pre), 2)
                warnings.append(f"Shot {sid} — Text End {shot['text_end']}s runs into the cut/transition; clamped to {new_end}s")
                shot["text_end"] = new_end
            if shot["text_end"] > d and kind in ("card", "logo_card"):
                shot["text_end"] = d
            if shot["text_end"] <= shot["text_start"]:
                warnings.append(f"Shot {sid} — Text End ≤ Text Start; overlay will be skipped")
            n_words = words(shot["text_overlay"])
            if kind == "card" and n_words > 8:
                warnings.append(f"Shot {sid} — card text is {n_words} words (README §2: ≤ 8). Trim the script line.")
            elif kind == "stock" and n_words > 4 and shot["text_style"] in ("auto", "keyword"):
                warnings.append(f"Shot {sid} — keyword overlay is {n_words} words (README §2: ≤ 4). "
                                f"Set Text Style = caption if it is a subtitle-style line.")
            if kind == "product":
                warnings.append(f"Shot {sid} — product insert carries a text overlay; README §2 says one element "
                                f"at a time. The overlay will be dropped at render.")
        if shot["transition_out"] in REAL_TRANSITIONS and shot["trans_dur"] <= 0:
            warnings.append(f"Shot {sid} — {shot['transition_out']} with Trans Dur 0; treated as a cut")
        if shot["trans_dur"] > d * 0.5:
            warnings.append(f"Shot {sid} — Trans Dur {shot['trans_dur']}s is over half the shot; clamped")
            shot["trans_dur"] = round(d * 0.5, 2)
        if shot["local_file"] and not Path(shot["local_file"]).exists() and shot["status"] == "downloaded":
            warnings.append(f"Shot {sid} — status 'downloaded' but file missing: {shot['local_file']}")
        if kind == "product" and not shot["local_file"]:
            errors.append(f"Shot {sid} — T4 product insert needs a Local File Path (screen recording)")
        shots.append(shot)

    return shots, warnings, errors


def restraint_audit(shots: list, settings: dict) -> list:
    """README §2 anti-clutter law."""
    notes = []
    overlay_budget = to_int(settings.get("overlay_budget"), 5)
    card_budget    = to_int(settings.get("card_budget"), 3)
    cards    = [s for s in shots if s["kind"] == "card"]
    keywords = [s for s in shots if s["kind"] == "stock" and s["text_overlay"]]
    products = [s for s in shots if s["kind"] == "product"]
    if len(cards) > card_budget:
        notes.append(f"Restraint: {len(cards)} blank text frames (T1) > budget {card_budget}")
    if len(keywords) > overlay_budget:
        notes.append(f"Restraint: {len(keywords)} stock overlays (T3) > budget {overlay_budget} — "
                     f"the weakest should be cut (README §2)")
    if len(products) > 3:
        notes.append(f"Restraint: {len(products)} product inserts (T4) > 3")
    # No two treatment events back-to-back; a card is followed by ~6s of clean T2
    tl = compute_timeline(shots)
    for i, s in enumerate(shots[:-1]):
        nxt = shots[i + 1]
        if s["kind"] == "card" and (nxt["kind"] == "card" or (nxt["kind"] == "stock" and nxt["text_overlay"])):
            notes.append(f"Restraint: shot {s['shot_id']} (card) is immediately followed by another "
                         f"treatment event ({nxt['shot_id']}); README §2 wants ~6s of clean footage between")
    orange_frames = [s for s in cards if hex_clean(s.get("card_bg") or "", "#000000") == BRAND["orange"]]
    if len(orange_frames) > 1:
        notes.append("Restraint: Brio Orange card used more than once (reserve for THE single most important frame)")
    total = timeline_total(tl)
    if total > to_float(settings.get("max_duration"), 90):
        notes.append(f"Length: assembled video is {total:.1f}s — over the {settings.get('max_duration', 90)}s Reels target")
    return notes


def carry_forward(shots: list, prev_path: Path | None) -> list:
    """
    Keep the assets the fetcher already chose (and the ones review rejected) when
    project.json is rebuilt from the sheet. A shot keeps its asset only while its
    brief (keywords, scene, asset priority) is unchanged, its file still exists,
    and the sheet neither names its own file nor asks for a swap.
    """
    notes: list = []
    if not prev_path or not prev_path.exists():
        return notes
    try:
        prev = json.loads(prev_path.read_text(encoding="utf-8-sig"))
    except Exception:
        return [f"Previous {prev_path.name} is unreadable — fetched assets are not carried forward"]
    prev_by_id = {str(s.get("shot_id")): s for s in prev.get("shots", [])}
    carried = 0
    for s in shots:
        p = prev_by_id.get(s["shot_id"])
        if not p or s["kind"] != "stock":
            continue
        rejected = list(dict.fromkeys(p.get("rejected_ids", [])))
        if s["local_file"]:
            if rejected:
                s["rejected_ids"] = rejected
            continue                                   # the sheet names a file — the sheet wins
        if s["status"] in ("swap", "error"):
            if p.get("asset_id"):
                rid = f"{p.get('source')}:{p.get('asset_type')}:{p.get('asset_id')}"
                if rid not in rejected:
                    rejected.append(rid)
            if rejected:
                s["rejected_ids"] = rejected
            continue                                   # the sheet asks for a different asset
        if rejected:
            s["rejected_ids"] = rejected
        same_brief = (p.get("keywords") == s["keywords"] and p.get("scene_desc") == s["scene_desc"]
                      and p.get("asset_priority") == s["asset_priority"])
        pf = p.get("local_file", "")
        if not (pf and Path(pf).exists() and p.get("status") == "downloaded"):
            continue
        if not same_brief:
            notes.append(f"Shot {s['shot_id']} — keywords/scene changed since the last fetch; a new asset will be fetched")
            continue
        for f in CARRY_FIELDS:
            if f in p:
                s[f] = p[f]
        carried += 1
    if carried:
        notes.append(f"info: kept {carried} previously fetched asset(s) from {prev_path.name} "
                     f"(set Status = swap on a shot to replace its asset)")
    return notes


# ── Build project.json ────────────────────────────────────────────────────────
def build_project(xlsx_path: Path, prev_json: Path | None = None) -> tuple[dict, list, list]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    for sheet in ("Project Settings", "Shot Plan"):
        if sheet not in wb.sheetnames:
            sys.exit(f"❌  '{sheet}' sheet not found in {xlsx_path.name}")

    settings = {**DEFAULT_SETTINGS, **{k: v for k, v in read_settings(wb).items() if v != ""}}
    subjects = [x.strip() for x in str(settings.get("illustration_only_subjects", "")).split(",") if x.strip()] or None
    shots, warnings, errors = read_shots(wb, subjects)
    carry_notes = carry_forward(shots, prev_json)
    for s in shots:
        for field, label in (("text_overlay", "on-screen text"), ("vo_line", "VO line")):
            st, hits = check_text(s.get(field, ""))
            if st == SAFE_FAIL:
                errors.append(f"Shot {s['shot_id']} — {label} is not suitable for a U audience: {'; '.join(hits)}")
            elif st != SAFE_PASS:
                warnings.append(f"Shot {s['shot_id']} — {label} needs a content-safety review: {'; '.join(hits)}")
    if not shots:
        errors.append("Shot Plan has no shots")
    warnings += restraint_audit(shots, settings)

    timeline = compute_timeline(shots)
    total_planned = round(sum(s["duration"] for s in shots), 2)
    project_root = xlsx_path.parent

    def sp(key, default_rel):
        v = _clean_path(settings.get(key, ""))
        return v if v else str(project_root / default_rel)

    audio_paths = {
        "music_path": _clean_path(settings.get("music_path", "")),
        "bgm_path":   _clean_path(settings.get("bgm_path", "")),
        "vo_path":    _clean_path(settings.get("vo_path", "")),
        "sting_path": _clean_path(settings.get("sting_path", "")),
    }
    for k, p in audio_paths.items():
        if p and not Path(p).exists():
            warnings.append(f"Audio: {k} not found — {p} (that layer will be skipped)")
    logo_path = _clean_path(settings.get("logo_path", ""))
    if logo_path and not Path(logo_path).exists():
        warnings.append(f"Logo: file not found — {logo_path}")
    if not audio_paths["vo_path"]:
        warnings.append("Audio: no vo_path — the video will have no voiceover (audio_mode B: music + captions)")

    # Only keys typed into the sheet are carried. Environment keys are read live by each stage
    # (env_key) and must never be copied into project.json, where they would sit on disk in plain text.
    api_keys = {
        "pexels":   _str(settings.get("pexels_api_key", "")),
        "pixabay":  _str(settings.get("pixabay_api_key", "")),
        "unsplash": _str(settings.get("unsplash_api_key", "")),
    }

    project = {
        "svos_version":   "2.0",
        "source_xlsx":    str(xlsx_path.resolve()),
        "project_name":   settings.get("project_name", "my_video"),
        "output_folder":  sp("output_folder", "Output"),
        "output_file":    settings.get("output_filename", "final_video.mp4"),
        "width":          to_int(settings.get("output_width"), 1080),
        "height":         to_int(settings.get("output_height"), 1920),
        "fps":            to_int(settings.get("fps"), 30),
        "platform":       settings.get("platform", "reels"),
        "total_duration": total_planned,
        "planned_assembled_duration": timeline_total(timeline),
        "audio": {
            **audio_paths,
            "music_volume":   to_float(settings.get("music_volume"), 0.8),
            "music_fade_in":  to_float(settings.get("music_fade_in"), 0.5),
            "music_fade_out": to_float(settings.get("music_fade_out"), 2.0),
            "bgm_volume":     to_float(settings.get("bgm_volume"), 0.4),
            "vo_volume":      to_float(settings.get("vo_volume"), 1.0),
            "vo_start":       to_float(settings.get("vo_start"), 0.0),
            "sting_volume":   to_float(settings.get("sting_volume"), 0.9),
            "target_lufs":    to_float(settings.get("target_lufs"), -14.0),
            "true_peak":      to_float(settings.get("true_peak"), -1.5),
            "duck_db":        to_float(settings.get("duck_db"), 8.0),
        },
        "logo": {
            "path":       logo_path,
            "position":   settings.get("logo_position", "bottom_right"),
            "scale":      to_float(settings.get("logo_scale"), 0.15),
            "opacity":    to_float(settings.get("logo_opacity"), 0.9),
            "mode":       settings.get("logo_mode", "end_only").lower(),
            "card_bg":    hex_clean(settings.get("logo_card_bg"), BRAND["purple"]),
            "card_scale": to_float(settings.get("logo_card_scale"), 0.42),
        },
        "api_keys":      api_keys,
        "assets_folder": sp("assets_folder", "Assets"),
        "render": {
            "crf":           str(settings.get("crf", "18")),
            "preset":        settings.get("preset", "slow"),
            "unified_grade": settings.get("unified_grade", "warm_soft").lower(),
            "film_grain":    to_bool(settings.get("film_grain"), False),
            "cover_time":    settings.get("cover_time", ""),
            "max_duration":  to_float(settings.get("max_duration"), 90.0),
        },
        "text": {
            "font":          settings.get("text_font", "Inter"),
            "font_tamil":    settings.get("text_font_tamil", ""),
            "fonts_folder":  _clean_path(settings.get("fonts_folder", "")),
            "size_default":  to_int(settings.get("text_size_default"), 72),
            "shadow":        to_bool(settings.get("text_shadow"), True),
            "box":           settings.get("text_box", "auto").lower(),
            "anim_default":  settings.get("text_anim_default", "rise").lower(),
            "brand_color":   hex_clean(settings.get("brand_color"), BRAND["orange"]),
            "card_bg":       hex_clean(settings.get("card_bg_default"), BRAND["offwhite"]),
            "card_text":     hex_clean(settings.get("card_text_color"), BRAND["near_black"]),
            "safe_zones":    to_bool(settings.get("safe_zones"), True),
            "overlay_budget": to_int(settings.get("overlay_budget"), 5),
            "card_budget":    to_int(settings.get("card_budget"), 3),
            "match_vo":       to_bool(settings.get("match_text_to_vo"), True),   # vo_aligner may fix mismatched text
            "cta_style":      str(settings.get("cta_style", "text") or "text").lower(),   # text | button
            "cta_color":      hex_clean(settings.get("cta_color"), "#B85F00"),           # button colour
            "cta_early":      to_bool(settings.get("cta_early"), True),                 # chip on the last shot
            "ab_variant":     to_bool(settings.get("ab_variant"), False),               # also render version B
        },
        "fetch": {
            "max_queries_per_shot": max(1, to_int(settings.get("max_queries_per_shot"), 3)),
            "budget": {p: to_int(settings.get(f"fetch_budget_{p}"), d)
                       for p, d in (("pexels", 150), ("pixabay", 300), ("unsplash", 40))},
            "cache_hours": to_float(settings.get("fetch_cache_hours"), 24.0),
            "max_download_mb": to_float(settings.get("max_download_mb"), 300.0),
        },
        "safety": {
            "vo_language": str(settings.get("vo_language", "auto")).lower(),
            "captions_path": _clean_path(settings.get("captions_path", "")),
            "models_dir": _clean_path(settings.get("safety_models_dir", "")),
            "frames_per_second": to_float(settings.get("safety_frames_per_second"), 1.0),
            "illustration_only_subjects": subjects or [],
        },
        "shots":    shots,
        "timeline": timeline,
        "validation": {"warnings": warnings, "errors": errors},
    }
    # T5 logo card: use the logo variant that is fully visible on this card's colour (a purple 'U'
    # vanishes on a purple card). Chosen here so the safety gate clears the exact file that is rendered.
    for s in shots:
        if shot_kind(s) == "logo_card":
            base = s.get("local_file") or logo_path
            bg = hex_clean(s.get("card_bg") or project["logo"]["card_bg"], BRAND["purple"])
            pick, note = best_logo(base, [bg]) if base and Path(base).exists() else (base, "")
            s["logo_file"] = pick
            if note:
                carry_notes.append(f"info: Shot {s['shot_id']} logo card on {bg} — {note}")
    if prev_json and prev_json.exists():
        try:   # safety decisions are keyed by content hash, so they stay valid when the sheet is re-read
            project["safety_state"] = json.loads(prev_json.read_text(encoding="utf-8-sig")).get("safety_state", {})
        except Exception:
            pass
    infos = [n[len("info: "):] for n in carry_notes if n.startswith("info: ")]
    warnings += [n for n in carry_notes if not n.startswith("info: ")]
    project["carried_forward"] = infos
    return project, warnings, errors


def main():
    parser = argparse.ArgumentParser(description="Parse story_plan.xlsx → project.json")
    parser.add_argument("xlsx", help="Path to story_plan.xlsx")
    parser.add_argument("--out", default=None, help="Output JSON path (default: project.json next to the xlsx)")
    parser.add_argument("--strict", action="store_true", help="Treat warnings as errors")
    parser.add_argument("--fresh", action="store_true",
                        help="Do not carry fetched assets forward from the existing project.json")
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    if not xlsx_path.exists():
        sys.exit(f"❌  File not found: {xlsx_path}")

    print(f"\n  Reading {xlsx_path.name}...")
    out_path = Path(args.out) if args.out else xlsx_path.parent / "project.json"
    project, warnings, errors = build_project(xlsx_path, None if args.fresh else out_path)
    out_path.write_text(json.dumps(project, indent=2, ensure_ascii=False), encoding="utf-8")

    kinds = {}
    for s in project["shots"]:
        kinds[s["treatment"]] = kinds.get(s["treatment"], 0) + 1
    print(f"  ✅  {len(project['shots'])} shots  |  planned {project['total_duration']}s  "
          f"|  assembled ≈ {project['planned_assembled_duration']}s (after transition overlaps)")
    print(f"  ✅  treatments: " + "  ".join(f"{k}×{v}" for k, v in sorted(kinds.items())))
    print(f"  ✅  project.json → {out_path.resolve()}")
    for note in project.get("carried_forward", []):
        print(f"  ✅  {note}")

    if warnings:
        print(f"\n  ⚠️  {len(warnings)} warning(s):")
        for w in warnings:
            print(f"     • {w}")
    if errors:
        print(f"\n  ❌  {len(errors)} error(s):")
        for e in errors:
            print(f"     • {e}")
        sys.exit(1)
    if args.strict and warnings:
        sys.exit("\n  ❌  --strict: warnings present")


if __name__ == "__main__":
    main()
